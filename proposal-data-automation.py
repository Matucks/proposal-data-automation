import pandas as pd
import openpyxl
import os
import glob
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders

# Caminhos das pastas e arquivos
INPUT_FOLDER = 'C:\\AutomationProject\\Input'
OUTPUT_PATH = 'C:\\AutomationProject\\Output\\Proposals.xlsx'
DATA_PATH = 'C:\\AutomationProject\\Data\\ActiveEmployees.xlsx'

# Buscar o arquivo Excel mais recente na pasta de entrada
input_files = glob.glob(os.path.join(INPUT_FOLDER, '*.xlsx'))
if not input_files:
    print(f"No Excel file found in the input folder: {INPUT_FOLDER}")
else:
    # Seleciona o arquivo Excel mais recente
    input_file_path = max(input_files, key=os.path.getctime)
    print(f"Input file found: {input_file_path}")

    # Carregar a planilha de entrada e o arquivo de dados de funcionários
    input_df = pd.read_excel(input_file_path)
    employees_df = pd.read_excel(DATA_PATH, sheet_name='Active')

    # Criar um dicionário para correspondência do vendedor com a unidade
    unit_lookup = employees_df.set_index('Name')['Unit'].to_dict()

    # Criar a aba DADOS BRUTOS
    raw_data_df = input_df.copy()

    # Adicionar a coluna 'Unit' baseada na coluna 'Seller', preenchendo com '#' se o vendedor não for encontrado
    raw_data_df['Unit'] = raw_data_df['Seller'].map(unit_lookup).fillna('#')

    # Remover colunas duplicadas, se existirem
    raw_data_df = raw_data_df.loc[:, ~raw_data_df.columns.duplicated()]

    # Reordenar as colunas para colocar 'Unit' ao lado de 'Seller'
    columns = list(raw_data_df.columns)
    seller_index = columns.index('Seller')
    columns.insert(seller_index + 1, columns.pop(columns.index('Unit')))
    raw_data_df = raw_data_df[columns]

    # Excluir a última linha
    raw_data_df = raw_data_df[:-1]

    # Remover duplicatas
    raw_data_df = raw_data_df.drop_duplicates()

    # Filtrar dados para Team A e Team B
    team_a_units = ["UnitA", "UnitB", "UnitC", "UnitD", "UnitE"]
    team_b_units = ["UnitF", "UnitG", "UnitH", "UnitI", "UnitJ"]

    team_a_df = raw_data_df[raw_data_df['Unit'].isin(team_a_units)]
    team_b_df = raw_data_df[raw_data_df['Unit'].isin(team_b_units)]

    # Criar tabelas dinâmicas separadas para cada filial
    units = raw_data_df['Unit'].unique()

    # Salvar os resultados em um arquivo Excel com abas separadas
    with pd.ExcelWriter(OUTPUT_PATH, engine='openpyxl') as writer:
        # Adicionar DADOS BRUTOS e criar tabela
        raw_data_df.to_excel(writer, sheet_name='Raw Data', index=False)
        workbook = writer.book
        raw_data_sheet = writer.sheets['Raw Data']
        raw_data_table = Table(displayName="RawDataTable", ref=f"A1:{chr(64 + len(raw_data_df.columns))}{len(raw_data_df) + 1}")
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=True
        )
        raw_data_table.tableStyleInfo = style
        raw_data_sheet.add_table(raw_data_table)

        # Adicionar Team A e criar tabela
        team_a_df.to_excel(writer, sheet_name='Team A', index=False)
        team_a_sheet = writer.sheets['Team A']
        team_a_table = Table(displayName="TeamATable", ref=f"A1:{chr(64 + len(team_a_df.columns))}{len(team_a_df) + 1}")
        team_a_table.tableStyleInfo = style
        team_a_sheet.add_table(team_a_table)

        # Adicionar Team B e criar tabela
        team_b_df.to_excel(writer, sheet_name='Team B', index=False)
        team_b_sheet = writer.sheets['Team B']
        team_b_table = Table(displayName="TeamBTable", ref=f"A1:{chr(64 + len(team_b_df.columns))}{len(team_b_df) + 1}")
        team_b_table.tableStyleInfo = style
        team_b_sheet.add_table(team_b_table)

        # Criar abas dinâmicas para cada filial
        for unit in units:
            unit_data = raw_data_df[raw_data_df['Unit'] == unit]
            unit_data.to_excel(writer, sheet_name=f"{unit}", index=False)
            unit_sheet = writer.sheets[unit]
            unit_table = Table(displayName=f"{unit}Table", ref=f"A1:{chr(64 + len(unit_data.columns))}{len(unit_data) + 1}")
            unit_table.tableStyleInfo = style
            unit_sheet.add_table(unit_table)

        # Criar tabela dinâmica consolidada
        pivot_table = pd.pivot_table(
            raw_data_df,
            values='Chassis',  # Supondo que Chassis seja a coluna de contagem
            index='Unit',
            columns='Status',
            aggfunc='count',
            margins=True,
            margins_name='Total'
        )
        pivot_table = pivot_table.fillna(0).astype(int).reset_index()

        pivot_table.to_excel(writer, sheet_name='Consolidated Pivot', index=False)
        pivot_sheet = writer.sheets['Consolidated Pivot']
        pivot_table_object = Table(displayName="ConsolidatedPivotTable", ref=f"A1:{chr(64 + len(pivot_table.columns))}{len(pivot_table) + 1}")
        pivot_table_object.tableStyleInfo = style
        pivot_sheet.add_table(pivot_table_object)

    print("Processing completed. File with Team A, Team B, unit-specific tables, and consolidated pivot table saved successfully.")

    # Enviar email com o arquivo gerado
    send_email_confirmation = input("File generated. Do you want to send it via email? (y/n): ").strip().lower()
    if send_email_confirmation == 'y':
        # Configurar informações do email
        sender_email = 'automation@example.com'
        sender_password = 'yourpassword'
        recipients = ['recipient1@example.com', 'recipient2@example.com']
        subject = 'Automated Proposal Report'
        body = (
            'Hello,

'
            'Please find attached the latest proposal report.

'
            'This email is sent automatically. If you encounter any issues, please contact the administrator.'
        )

        # Configurar a mensagem do email
        message = MIMEMultipart()
        message['From'] = sender_email
        message['To'] = ', '.join(recipients)
        message['Subject'] = subject
        message.attach(MIMEText(body, 'plain'))

        # Anexar o arquivo gerado
        with open(OUTPUT_PATH, 'rb') as attachment:
            mime_base = MIMEBase('application', 'octet-stream')
            mime_base.set_payload(attachment.read())
        encoders.encode_base64(mime_base)
        mime_base.add_header('Content-Disposition', f'attachment; filename={os.path.basename(OUTPUT_PATH)}')
        message.attach(mime_base)

        # Enviar o email
        try:
            server = smtplib.SMTP('smtp.example.com', 587)
            server.starttls()
            server.login(sender_email, sender_password)
            server.sendmail(sender_email, recipients, message.as_string())
            server.quit()
            print("Email sent successfully!")
        except Exception as e:
            print(f"Failed to send email: {e}")
    else:
        print("Email sending canceled.")
